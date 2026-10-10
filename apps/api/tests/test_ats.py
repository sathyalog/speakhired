import io
import os
import tempfile

os.environ.setdefault("DATA_DIR", tempfile.mkdtemp())

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402

from speakhired import main  # noqa: E402
from speakhired.agents import resume as R  # noqa: E402
from speakhired.config import Settings  # noqa: E402
from speakhired.engines import ats  # noqa: E402
from speakhired.ingest.documents import inspect_resume  # noqa: E402
from speakhired.reasoning.lev import LevClient, LevUnavailable  # noqa: E402

GOOD = """Priya Sharma
priya.sharma@example.com | +91 98765 43210 | linkedin.com/in/priya-sharma | Hyderabad

Summary
Backend engineer with 4 years of experience building Kafka and Python services for payments.

Skills
Python, Java, Kafka, PostgreSQL, Docker, Kubernetes, REST APIs, CI/CD

Work Experience
Acme Pay, Software Engineer, Jan 2021 - Present
• Built a Kafka consumer service that processed 2 million events a day with 99.9% uptime
• Reduced API latency by 40% by adding Redis caching to the orders service
• Led a team of 4 engineers to migrate 12 services to Kubernetes in 3 months
• Automated deployments with Jenkins, cutting release time from 2 hours to 15 minutes
• Designed PostgreSQL schemas for 3 new products and mentored 2 junior developers

Education
B.Tech in Computer Science, JNTU Hyderabad, 2017 - 2021
"""


def docx_bytes(text: str, table: bool = False, header: str = "") -> bytes:
    d = Document()
    if header:
        d.sections[0].header.paragraphs[0].text = header
    if table:
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = text[: len(text) // 2], text[len(text) // 2:]
    else:
        for line in text.splitlines():
            d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pdf_bytes(lines: list[str], two_columns: bool = False, blank: bool = False) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    if blank:
        c.rect(50, 50, 400, 300)
    else:
        y = 800
        for ln in lines:
            c.drawString(50, y, ln)
            if two_columns:
                c.drawString(330, y, "Skills Python Java Kafka")
            y -= 14
            if y < 40:
                break
    c.save()
    return buf.getvalue()


class FakeLLM:
    def __init__(self, fail=False):
        self.fail, self.calls = fail, []

    async def chat_json(self, system, user, schema, temperature=0.2):
        self.calls.append(user)
        if self.fail:
            raise RuntimeError("offline")
        props = schema.get("properties", {})
        if "has_summary" in props:
            return {"has_summary": True, "action_verbs": 4, "quantified": 4, "clarity": 4, "specific_tech": 4}
        if "keywords" in props:
            return {"keywords": ["Python", "Kafka", "GraphQL", "k8s"]}
        if "rewrites" in props:
            return {"rewrites": [{"before": "Responsible for fixing bugs in the billing module", "after": "Fixed [add number] billing bugs, improving checkout reliability"},
                                 {"before": "A line that is not in the resume", "after": "x"},
                                 {"before": "Worked on login page for users", "after": "Rebuilt the login page for 5000 users"}]}
        return {}

    async def chat_text(self, system, user, temperature=0.6):
        return "ok"


class SpyLev(LevClient):
    def __init__(self, settings):
        super().__init__(settings)
        self.states = []

    async def decide(self, state, questions):
        self.states.append(state)
        return {"has_summary": {"noul": 0.9}, **{k: {"score": 3.2, "confidence": 0.9} for k in ("action_verbs", "quantified", "clarity", "specific_tech")}}


REMOTE = Settings(lev_mode="http", lev_url="https://me--lev.modal.run", lev_allow_remote=True)
LOCAL = Settings(lev_mode="http", lev_url="http://127.0.0.1:8000")


def facts(text=GOOD, **kw):
    return inspect_resume("r.docx", docx_bytes(text, **kw))


async def run(f, jd="", llm=None, lev=None, use_lev=False, rewrites=False):
    return await R.analyze(llm or FakeLLM(), lev or LevClient(Settings(lev_mode="off")), f, jd, use_lev, rewrites)


async def test_good_resume_scores_high_and_is_explained():
    r = await run(facts())
    assert r.score >= 80 and r.grade in ("Good", "Excellent") and r.judge == "Ollama"
    assert {s.id for s in r.sections} == {"format", "structure", "content", "readability"}  # keywords left out without a JD
    assert any("job description" in n for n in r.notes)


async def test_table_layout_and_missing_contact_score_lower_with_fixes():
    good = await run(facts())
    bad = await run(facts(GOOD.replace("priya.sharma@example.com", "").replace("+91 98765 43210", ""), table=True))
    assert bad.score < good.score and bad.score <= 69 and bad.grade in ("Fair", "Needs work")
    assert any("capped" in n for n in bad.notes)
    failed = {c.id for c in bad.checks if not c.passed}
    assert {"tables", "email", "phone"} <= failed
    titles = [s.title for s in bad.suggestions]
    assert "Remove tables" in titles and "Add a phone number" in titles
    assert bad.suggestions[0].priority == "high"


async def test_score_rises_when_problems_are_fixed():
    base = GOOD.replace("linkedin.com/in/priya-sharma", "").replace("Summary", "About")
    fixed = await run(facts(GOOD))
    worse = await run(facts(base))
    assert fixed.score > worse.score


def test_scanned_pdf_is_flagged_unreadable():
    f = inspect_resume("scan.pdf", pdf_bytes([], blank=True))
    assert f.extractable is False
    checks = ats.run_checks(f)
    assert not next(c for c in checks if c.id == "extractable").passed


async def test_unreadable_resume_gets_low_score_and_no_ai_content():
    r = await run(inspect_resume("scan.pdf", pdf_bytes([], blank=True)))
    assert r.score < 30 and {s.id for s in r.sections} <= {"format", "structure", "readability"}


def test_pdf_columns_heuristic():
    lines = [f"Line {i} some resume text here" for i in range(30)]
    assert inspect_resume("a.pdf", pdf_bytes(lines)).columns_likely is False
    assert inspect_resume("b.pdf", pdf_bytes(lines, two_columns=True)).columns_likely is True
    assert inspect_resume("a.pdf", pdf_bytes(GOOD.splitlines() * 2)).extractable is True


def test_docx_header_contact_is_not_enough():
    f = facts(GOOD.replace("priya.sharma@example.com", ""), header="priya.sharma@example.com")
    email = next(c for c in ats.run_checks(f) if c.id == "email")
    assert not email.passed and 0 < email.credit < 1


def test_keyword_matching_with_aliases():
    kw = ats.match_keywords("Built APIs in JS and Postgres on k8s, CI/CD, Java 17", ["JavaScript", "PostgreSQL", "Kubernetes", "Java", "C++", "Go", "GraphQL", "continuous integration"])
    assert kw.matched == ["JavaScript", "PostgreSQL", "Kubernetes", "Java", "continuous integration"] and kw.missing == ["C++", "Go", "GraphQL"]
    assert ats.match_keywords("I know javascript only", ["Java"]).missing == ["Java"]  # no partial-word match
    assert "kubernetes" in ats.vocabulary_keywords("Experience with K8s required")


async def test_keywords_section_added_with_a_job_description():
    r = await run(facts(), jd="We need Python, Kafka, GraphQL and k8s experience.")
    assert "keywords" in {s.id for s in r.sections}
    assert "GraphQL" in r.keywords.missing and "Python" in r.keywords.matched and "k8s" in r.keywords.matched
    assert any(s.category == "Keywords" for s in r.suggestions)


def test_remote_copy_is_redacted_and_trimmed():
    out = R.prepare_for_remote(GOOD)
    for secret in ("priya.sharma@example.com", "98765", "linkedin.com", "Priya Sharma"):
        assert secret not in out
    assert "Kafka" in out and len(out) <= 5000


async def test_remote_lev_needs_opt_in():
    lev = SpyLev(REMOTE)
    off = await run(facts(), lev=lev, use_lev=False)
    assert lev.states == [] and off.judge == "Ollama" and any("remote server" in n for n in off.notes)
    on = await run(facts(), lev=lev, use_lev=True)
    assert len(lev.states) == 1 and on.judge == "LEV" and "priya.sharma@example.com" not in lev.states[0]


async def test_local_lev_used_by_default():
    lev = SpyLev(LOCAL)
    r = await run(facts(), lev=lev)
    assert r.judge == "LEV" and len(lev.states) == 1


async def test_lev_failure_falls_back_to_ollama():
    class Down(SpyLev):
        async def decide(self, state, questions):
            raise LevUnavailable("down")
    r = await run(facts(), lev=Down(LOCAL))
    assert r.judge == "Ollama" and any("unavailable" in n for n in r.notes)


async def test_rewrites_keep_only_real_quotes_without_invented_numbers():
    text = GOOD + "\n• Responsible for fixing bugs in the billing module\n• Worked on login page for users\n"
    r = await run(facts(text), rewrites=True)
    rewrites = [s for s in r.suggestions if s.category == "Bullets"]
    assert [s.before for s in rewrites] == ["Responsible for fixing bugs in the billing module"]  # fake quote and invented 5000 dropped
    assert "[add number]" in rewrites[0].after


async def test_model_offline_still_returns_a_report():
    r = await run(facts(), llm=FakeLLM(fail=True))
    assert r.judge == "rules only" and r.score > 0 and any("could not be reached" in n for n in r.notes)


def test_endpoint_upload_limits_and_success():
    main.llm = FakeLLM()
    main.lev = LevClient(Settings(lev_mode="off"))
    c = TestClient(main.app)
    ok = c.post("/api/ats/analyze", files={"resume": ("cv.docx", docx_bytes(GOOD), "application/octet-stream")}, data={"jd_text": "Python and Kafka"})
    assert ok.status_code == 200 and ok.json()["score"] > 50 and "sections" in ok.json()
    assert c.post("/api/ats/analyze", files={"resume": ("cv.exe", b"x", "application/octet-stream")}).status_code == 400
    assert c.post("/api/ats/analyze", files={"resume": ("cv.pdf", b"not a pdf", "application/pdf")}).status_code == 400
    assert c.post("/api/ats/analyze", files={"resume": ("big.pdf", b"x" * (5 * 1024 * 1024 + 1), "application/pdf")}).status_code == 413


def test_year_ranges_are_not_phone_numbers():
    from speakhired.privacy import has_phone, redact
    assert not has_phone("B.Tech, JNTU Hyderabad, 2017 - 2021 and Jan 2021 - Present")
    assert has_phone("Call +91 98765 43210") and has_phone("9876543210")
    assert "2017 - 2021" in redact("Studied 2017 - 2021 at JNTU. Call +91 98765 43210.") and "98765" not in redact("Call +91 98765 43210")


async def test_critical_problems_cap_the_score():
    ok = await run(facts())
    assert not any("capped" in n for n in ok.notes)
    no_email = await run(facts(GOOD.replace("priya.sharma@example.com", "")))
    assert no_email.score <= 69 and any("no email" in n for n in no_email.notes)
    table = await run(facts(table=True))
    assert table.score <= 74 and any("table" in n for n in table.notes)


def test_odd_pdf_bullet_glyphs_are_recognised_and_date_mix_is_fine():
    f = inspect_resume("a.pdf", pdf_bytes(["Experience", "\x7f Built a Kafka service for payments with 2 million events", "\uf0b7 Reduced latency by 40% using Redis caching today", "Jan 2021 - Present", "2017 - 2021"]))
    assert len(ats.bullets_of(f.text)) == 2
    assert next(c for c in ats.run_checks(inspect_resume("a.docx", docx_bytes(GOOD))) if c.id == "dates").passed
    mixed = ats.run_checks(facts("Experience\nJan 2020 - Mar 2021\n03/2021 - 05/2022\n04/2022 - 06/2023\n"))
    assert not next(c for c in mixed if c.id == "dates").passed


def test_keyword_ai_call_skipped_when_vocabulary_is_enough():
    import asyncio
    llm = FakeLLM()
    kws = asyncio.run(R.extract_keywords(llm, "Python, Java, Kafka, Docker, Kubernetes, PostgreSQL, Redis and AWS required."))
    assert len(kws) >= 6 and llm.calls == []


def test_rewrites_endpoint_is_separate():
    main.llm = FakeLLM()
    c = TestClient(main.app)
    text = GOOD + "\n• Responsible for fixing bugs in the billing module\n"
    r = c.post("/api/ats/rewrites", files={"resume": ("cv.docx", docx_bytes(text), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["suggestions"][0]["before"].startswith("Responsible for fixing")


async def test_instant_mode_skips_all_ai_calls():
    llm = FakeLLM()
    r = await R.analyze(llm, LevClient(Settings(lev_mode="off")), facts(), "Python, Java, Kafka, Docker, Kubernetes, PostgreSQL", False, use_ai=False)
    assert llm.calls == [] and r.judge == "rules only" and r.score > 50 and any("skipped" in n for n in r.notes)


def test_judge_input_is_condensed():
    c = R.condense_for_judge(GOOD)
    assert "Built a Kafka consumer" in c and len(c) < len(GOOD)


class CountingLLM(FakeLLM):
    async def chat_json(self, system, user, schema, temperature=0.2):
        return await super().chat_json(system, user, schema, temperature)


async def test_online_mode_asks_the_language_model_even_when_lev_answers():
    lev, llm = SpyLev(LOCAL), CountingLLM()
    online = await R.analyze(llm, lev, facts(), "", False, online=True)
    assert online.judge == "LEV + Grok" and any("has_summary" in str(c) or c for c in llm.calls) and len(lev.states) == 1
    llm2 = CountingLLM()
    offline = await R.analyze(llm2, SpyLev(LOCAL), facts(), "", False, online=False)
    assert offline.judge == "LEV" and llm2.calls == []  # offline: Ollama is only a fallback, which keeps small computers fast


async def test_online_mode_keeps_going_when_grok_is_unavailable():
    from speakhired.llm import LLMUnavailable

    class NoGrok(FakeLLM):
        async def chat_json(self, system, user, schema, temperature=0.2):
            raise LLMUnavailable("grok_not_configured", "Grok key is not configured.")
    r = await R.analyze(NoGrok(), SpyLev(LOCAL), facts(), "", False, online=True)
    assert r.judge == "LEV" and any("Grok key is not configured" in n for n in r.notes)
    r2 = await R.analyze(NoGrok(), LevClient(Settings(lev_mode="off")), facts(), "", False, online=True)
    assert r2.judge == "rules only"
