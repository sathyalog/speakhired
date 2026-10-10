import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError(f"unsupported file type: {ext}")


_Q_LINE = re.compile(r"^\s*(?:\d+[.)]|[-*•Q]\d*[.):]?)\s*(.+\?)\s*$", re.I)


def split_questions(text: str) -> list[str]:
    """Pull question-like lines from an uploaded question bank (numbered, bulleted or plain '?' lines)."""
    seen, out = set(), []
    for line in text.splitlines():
        m = _Q_LINE.match(line) or (re.match(r"^\s*(.{12,}\?)\s*$", line))
        if m:
            q = m.group(1).strip()
            key = q.lower()
            if key not in seen:
                seen.add(key)
                out.append(q)
    return out


_LEAD_GLYPH = re.compile(r"^[ \t]*[\x00-\x08\x0b-\x1f\x7f\uf000-\uf8ff\u25a0-\u25ff\u2700-\u27bf\u2043\u2219\u00b7\u00a7]\s*", re.M)


def normalise_resume_text(text: str) -> str:
    """PDF and Word files often encode bullets as private-use or control characters; make them a plain bullet."""
    text = _LEAD_GLYPH.sub("• ", text)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)


@dataclass
class ResumeFacts:
    text: str
    file_type: str
    pages: int = 1
    tables: int = 0
    images: int = 0
    text_boxes: int = 0
    columns_likely: bool = False
    header_footer_text: str = ""
    extractable: bool = True
    notes: list[str] = field(default_factory=list)


def _pdf_columns(reader: PdfReader) -> bool:
    """Heuristic: many text rows that have words in both the left and the right part of the page."""
    rows_with_both = 0
    for page in reader.pages[:2]:
        width = float(page.mediabox.width) or 600.0
        rows: dict[int, list[float]] = {}

        def visit(text, cm, tm, font_dict, font_size):
            if text and len(text.strip()) >= 3:
                rows.setdefault(round(tm[5] / 3), []).append(tm[4] / width)

        try:
            page.extract_text(visitor_text=visit)
        except Exception:  # noqa: BLE001 - layout hints are best-effort
            return False
        rows_with_both += sum(1 for xs in rows.values() if min(xs) < 0.3 and max(xs) > 0.5 and len(xs) >= 2)
    return rows_with_both >= 8


def inspect_resume(filename: str, data: bytes) -> ResumeFacts:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        text = normalise_resume_text("\n".join((p.extract_text() or "") for p in reader.pages).strip())
        images = sum(len(list(p.images)) for p in reader.pages) if reader.pages else 0
        return ResumeFacts(text=text, file_type="pdf", pages=len(reader.pages), images=images,
                           columns_likely=_pdf_columns(reader), extractable=len(text) >= 150 * max(1, min(len(reader.pages), 2)) // 2)
    if ext == "docx":
        doc = Document(io.BytesIO(data))
        paras = [p.text for p in doc.paragraphs]
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        head = " ".join(p.text for sec in doc.sections for part in (sec.header, sec.footer) for p in part.paragraphs)
        text = normalise_resume_text("\n".join(paras + cells).strip())
        words = len(text.split())
        return ResumeFacts(text=text, file_type="docx", pages=max(1, round(words / 450)), tables=len(doc.tables),
                           images=len(doc.inline_shapes), text_boxes=doc.element.body.xml.count("txbxContent"),
                           header_footer_text=head, extractable=bool(text))
    if ext in ("txt", "md"):
        text = data.decode("utf-8", errors="ignore").strip()
        return ResumeFacts(text=text, file_type="txt", pages=max(1, round(len(text.split()) / 450)), extractable=bool(text))
    raise ValueError("Please upload a PDF, Word (.docx) or text resume.")
